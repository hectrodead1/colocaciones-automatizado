import openpyxl
from pathlib import Path
from datetime import datetime

class ProcesadorBanreservas:
    def procesar(self, filepath: str, datos_validacion: dict):
        """
        Lee el archivo de entrada Excel y genera una lista de líneas de texto
        para el archivo de texto de salida de Banreservas.
        
        Retorna (lineas_texto, lista_prestamos)
        """
        wb_in = openpyxl.load_workbook(filepath, data_only=True)
        ws_in = wb_in.active
        
        col_indices = datos_validacion['col_indices']
        start_row = datos_validacion['data_start_row']
        end_row = datos_validacion['data_end_row']
        
        lineas_texto = []
        lista_prestamos = []
        
        for in_row in range(start_row, end_row + 1):
            # Extraer valores
            cta_debitar = str(ws_in.cell(row=in_row, column=col_indices['CTA. A DEBITAR']).value or "").strip()
            cta_acreditar = str(ws_in.cell(row=in_row, column=col_indices['CTA. A CREDITAR']).value or "").strip()
            desembolso = ws_in.cell(row=in_row, column=col_indices['DESEMBOLSO']).value
            no_desembolso = str(ws_in.cell(row=in_row, column=col_indices['NO. DESEMBOLSO']).value or "").strip()
            
            # Nombre y PST para el registro mensual histórico
            nombre_in = str(ws_in.cell(row=in_row, column=col_indices['NOMBRE**']).value or "").strip()
            pst_in = str(ws_in.cell(row=in_row, column=col_indices['NUMERO PST']).value or "").strip()
            
            lista_prestamos.append({
                'numero_pst': pst_in,
                'nombre': nombre_in
            })
            
            # 1. Cuenta a debitar: sin guiones ni espacios, con 10001 delante, después una coma
            cta_debitar_limpia = cta_debitar.replace("-", "").replace(" ", "")
            if not cta_debitar_limpia.startswith("10001"):
                cta_debitar_final = f"10001{cta_debitar_limpia}"
            else:
                cta_debitar_final = cta_debitar_limpia
            parte_debitar = f"{cta_debitar_final},"
            
            # 2. Cuenta a acreditar: sin guiones ni espacios, con 20001 delante, después otra coma
            cta_acreditar_limpia = cta_acreditar.replace("-", "").replace(" ", "")
            if not cta_acreditar_limpia.startswith("20001"):
                cta_acreditar_final = f"20001{cta_acreditar_limpia}"
            else:
                cta_acreditar_final = cta_acreditar_limpia
            parte_acreditar = f"{cta_acreditar_final},"
            
            # 3. Desembolso: sin comas después otra coma
            if isinstance(desembolso, (int, float)):
                # Si es float/int, representarlo como número con decimales sin separador de miles
                # Por ejemplo: 46232.33
                monto_str = f"{desembolso:.2f}"
            else:
                # Si viene como string, le quitamos las comas
                monto_str = str(desembolso or "0.00").replace(",", "").strip()
            
            parte_desembolso = f"{monto_str},"
            
            # 4. Número de desembolso: eliminando los tres ceros
            # Si el valor empieza con "000", le quitamos los primeros 3 caracteres
            if no_desembolso.startswith("000"):
                parte_no_desembolso = no_desembolso[3:]
            else:
                parte_no_desembolso = no_desembolso.lstrip("0")
                if not parte_no_desembolso:  # Si quedara vacío
                    parte_no_desembolso = "0"
            
            # Unir toda la línea
            linea = f"{parte_debitar}{parte_acreditar}{parte_desembolso}{parte_no_desembolso}"
            lineas_texto.append(linea)
            
        wb_in.close()
        return lineas_texto, lista_prestamos
