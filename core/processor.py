import openpyxl
from openpyxl.styles import Font, Alignment, Protection
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.protection import SheetProtection
from openpyxl.worksheet.page import PageMargins
from datetime import datetime
from config import (
    COLUMNAS_SALIDA,
    TIPO_IDENTIFICACION,
    CONCEPTO,
    TIPO_REGISTRO,
    CORREO_BENEFICIARIO,
    ANCHOS_COLUMNA_SALIDA,
    ALTURA_FILA_DATOS,
    PLANTILLA_EXCEL_PATH,
    COLUMNAS_BLOQUEADAS,
    COLUMNAS_EDITABLES,
)


class Procesador:
    def procesar(self, filepath: str, datos_validacion: dict):
        """
        Lee el archivo de entrada y genera el workbook de salida para BHD.
        - Mantiene el Excel limpio sin colores extras
        - Bloquea solo las columnas B, C, D, E (Cédula, Beneficiario, Monto, Número único)
        - Protege la hoja sin contraseña para permitir fácil desprotección si es necesario
        """
        wb_in = openpyxl.load_workbook(filepath, data_only=True)
        ws_in = wb_in.active

        col_indices = datos_validacion['col_indices']
        start_row  = datos_validacion['data_start_row']
        end_row    = datos_validacion['data_end_row']

        # -----------------------------------------------------------------
        # Cargar plantilla especial o crear libro en blanco
        # -----------------------------------------------------------------
        if PLANTILLA_EXCEL_PATH.exists():
            wb_out = openpyxl.load_workbook(PLANTILLA_EXCEL_PATH)
            ws_out = wb_out.active
        else:
            wb_out = openpyxl.Workbook()
            ws_out = wb_out.active
            ws_out.title = "Hoja1"

        lista_prestamos = []

        # -----------------------------------------------------------------
        # Estilos reutilizables (sin relleno de color)
        # -----------------------------------------------------------------
        header_font  = Font(name='Calibri', size=11, bold=True, color=Color(theme=1))
        header_align = Alignment(horizontal='left', vertical='center')

        data_font_tnr       = Font(name='Times New Roman', size=10, color=Color(rgb='FF080000'))
        data_font_cal_11    = Font(name='Calibri', size=11, color=Color(theme=1))
        data_font_cal_12    = Font(name='Calibri', size=12, color=Color(rgb='FF000000'))

        align_left   = Alignment(horizontal='left',   vertical='center')

        # -----------------------------------------------------------------
        # Paso 1 – Escribir encabezados (solo si no hay plantilla)
        # -----------------------------------------------------------------
        if not PLANTILLA_EXCEL_PATH.exists():
            for col_idx, nombre_col in enumerate(COLUMNAS_SALIDA, 1):
                cell = ws_out.cell(row=1, column=col_idx, value=nombre_col)
                cell.font      = header_font
                cell.alignment = header_align
                cell.number_format = '@'

        # -----------------------------------------------------------------
        # Paso 2 – Escribir datos fila por fila
        # -----------------------------------------------------------------
        fecha_hoy = datetime.now().strftime("%d/%m/%Y")

        for in_row in range(start_row, end_row + 1):
            out_row = in_row - start_row + 2

            # Extraer valores desde el archivo de entrada
            cedula_in = str(ws_in.cell(row=in_row, column=col_indices['CEDULA']).value or "")
            nombre_in = str(ws_in.cell(row=in_row, column=col_indices['NOMBRE**']).value or "")
            monto_in  = ws_in.cell(row=in_row, column=col_indices['DESEMBOLSO']).value
            pst_in    = str(ws_in.cell(row=in_row, column=col_indices['NUMERO PST']).value or "")

            # Limpiar cédula (quitar guiones)
            cedula_limpia = cedula_in.replace("-", "").strip()

            # Registrar para el log mensual
            lista_prestamos.append({'numero_pst': pst_in, 'nombre': nombre_in})

            # ---- Col A: Tipo de identificación ----
            cA = ws_out.cell(row=out_row, column=1, value=TIPO_IDENTIFICACION)
            cA.font        = data_font_cal_11
            cA.alignment   = align_left
            cA.number_format = '@'

            # ---- Col C: Beneficiario ----
            cC = ws_out.cell(row=out_row, column=3, value=nombre_in)
            cC.font        = data_font_tnr
            cC.number_format = '@'

            # ---- Col D: Monto ----
            cD = ws_out.cell(row=out_row, column=4, value=float(monto_in) if monto_in else 0)
            cD.font        = data_font_tnr
            cD.number_format = '###,###,##0.00;\(###,###,##0.00\)'

            # ---- Col E: Número único ----
            cE = ws_out.cell(row=out_row, column=5, value=pst_in)
            cE.font        = data_font_tnr
            cE.number_format = '@'

            # ---- Col F: Fecha de efectividad ----
            cF = ws_out.cell(row=out_row, column=6, value=fecha_hoy)
            cF.font        = data_font_cal_11
            cF.alignment   = align_left
            cF.number_format = '@'

            # ---- Col G: Concepto ----
            cG = ws_out.cell(row=out_row, column=7, value=CONCEPTO)
            cG.font        = data_font_cal_11
            cG.number_format = '@'

            # ---- Col H: Tipo de registro ----
            cH = ws_out.cell(row=out_row, column=8, value=TIPO_REGISTRO)
            cH.font        = data_font_cal_11
            cH.alignment   = align_left
            cH.number_format = '@'

            # ---- Col I: Correo del beneficiario ----
            cI = ws_out.cell(row=out_row, column=9, value=CORREO_BENEFICIARIO)
            cI.font        = data_font_cal_12
            cI.number_format = 'General'

            # ---- Col B: Número de identificación (AL FINAL por requerimiento) ----
            cB = ws_out.cell(row=out_row, column=2, value=cedula_limpia)
            cB.font        = data_font_tnr
            cB.number_format = '@'

            # Altura de fila
            ws_out.row_dimensions[out_row].height = ALTURA_FILA_DATOS

        # -----------------------------------------------------------------
        # Paso 3 – Aplicar anchos de columna
        # -----------------------------------------------------------------
        for col_letter, width in ANCHOS_COLUMNA_SALIDA.items():
            ws_out.column_dimensions[col_letter].width = width

        # -----------------------------------------------------------------
        # Paso 4 – Bloquear únicamente las columnas B, C, D, E y proteger la hoja sin contraseña
        # -----------------------------------------------------------------
        self._aplicar_proteccion(ws_out, end_row - start_row + 1)

        # Establecer área de impresión de A a I
        ws_out.print_area = f"A1:I{out_row}"

        # Orientación horizontal y ajustar hoja a una página
        ws_out.page_setup.orientation = 'landscape'
        ws_out.page_setup.fitToPage = True
        ws_out.page_setup.fitToWidth = 1
        ws_out.page_setup.fitToHeight = 1
        ws_out.sheet_properties.pageSetUpPr.fitToPage = True

        wb_in.close()
        return wb_out, lista_prestamos

    def _aplicar_proteccion(self, ws, num_filas_datos: int):
        """
        Bloquea únicamente las celdas de las columnas B, C, D, E y encabezados.
        Deja desbloqueadas las columnas A, F, G, H, I.
        Protege la hoja sin contraseña.
        """
        total_cols = len(COLUMNAS_SALIDA)
        fila_header = 1
        fila_inicio = 2
        fila_fin    = fila_inicio + num_filas_datos - 1

        # --- Paso A: desbloquear todas las celdas de datos por defecto ---
        for fila in range(fila_inicio, fila_fin + 1):
            for col in range(1, total_cols + 1):
                ws.cell(row=fila, column=col).protection = Protection(locked=False)

        # --- Paso B: bloquear únicamente las columnas B, C, D, E ---
        for fila in range(fila_inicio, fila_fin + 1):
            for col_letra in COLUMNAS_BLOQUEADAS:
                col_idx = column_index_from_string(col_letra)
                ws.cell(row=fila, column=col_idx).protection = Protection(locked=True)

        # --- Paso C: bloquear encabezados ---
        for col in range(1, total_cols + 1):
            ws.cell(row=fila_header, column=col).protection = Protection(locked=True)

        # --- Paso D: activar protección de hoja sin contraseña ---
        ws.protection = SheetProtection(
            sheet               = True,
            selectLockedCells   = True,
            selectUnlockedCells = True,
            formatCells         = False,
            insertRows          = False,
            deleteRows          = False,
            sort                = False,
            autoFilter          = False,
        )
