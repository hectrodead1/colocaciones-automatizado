import openpyxl
from pathlib import Path
from datetime import datetime
from config import CARPETA_REGISTROS, COLUMNAS_REGISTRO, MESES_ES

class RegistroMensual:
    def registrar(self, lista_prestamos: list, nombre_archivo_generado: str):
        """
        Añade los préstamos al registro mensual correspondiente.
        Nunca borra registros anteriores.
        """
        ahora = datetime.now()
        año_actual = str(ahora.year)
        mes_actual_es = MESES_ES[ahora.month]
        
        # Estructura: Control de Préstamos / YYYY / Mes
        carpeta_mes = CARPETA_REGISTROS / año_actual / mes_actual_es
        carpeta_mes.mkdir(parents=True, exist_ok=True)
        
        nombre_registro = f"Registro {mes_actual_es} {año_actual}.xlsx"
        ruta_registro = carpeta_mes / nombre_registro
        
        if ruta_registro.exists():
            wb = openpyxl.load_workbook(ruta_registro)
            ws = wb.active
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Registro"
            # Escribir headers
            for col, header in enumerate(COLUMNAS_REGISTRO, 1):
                ws.cell(row=1, column=col, value=header)
                
        # Encontrar la primera fila vacía
        next_row = ws.max_row + 1
        
        fecha_str = ahora.strftime("%d/%m/%Y")
        hora_str = ahora.strftime("%H:%M")
        
        # Añadir registros
        for prestamo in lista_prestamos:
            ws.cell(row=next_row, column=1, value=fecha_str)
            ws.cell(row=next_row, column=2, value=hora_str)
            ws.cell(row=next_row, column=3, value=prestamo['numero_pst'])
            ws.cell(row=next_row, column=4, value=prestamo['nombre'])
            ws.cell(row=next_row, column=5, value=nombre_archivo_generado)
            next_row += 1
            
        wb.save(ruta_registro)
        wb.close()
